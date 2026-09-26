"""Модульные тесты: шаблоны, служебные слайды, предпросмотр, таблица, проверка, история."""

import io
import zipfile

import pytest
from lxml import etree
from pptx import Presentation

import extras
import generator as g
import preview
import review
from history import History


def question(n, text="Кто написал «Войну и мир»?", answer="Лев Толстой", **extra):
    q = {"number": n, "orig_number": n, "question": text, "answer": answer, "accept": None,
         "comment": None, "hard": False, "tour": None, "q_pictures": [], "a_pictures": []}
    q.update(extra)
    return q


def custom_template(with_notes: bool = False) -> bytes:
    prs = Presentation()
    for body in ("Текст вопроса", "Ответ"):
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = "Вопрос 1"
        slide.placeholders[1].text = body
        if with_notes:
            slide.notes_slide.notes_text_frame.text = "старая заметка шаблона"
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


# ─── Генератор ───────────────────────────────────────────────────────────────

def test_service_slides_and_tour_separators():
    qs = [question(1, tour="Тур 1"), question(2, tour="Тур 1"), question(3, tour="Тур 2")]
    specs = g.build_specs(qs, {"service": True}, g.standard_template(), "Кубок")
    assert [s.kind for s in specs] == ["title", "tour", "question", "answer", "question", "answer",
                                       "tour", "question", "answer", "final"]
    assert specs[0].label is None and specs[0].anchor == "ctr"
    # С перемешиванием разделители туров не нужны
    shuffled = g.build_specs(qs, {"service": True, "shuffle": True}, g.standard_template(), "Кубок")
    assert [s.kind for s in shuffled].count("tour") == 0


def test_notes_and_question_on_answer_slide():
    data = g.generate_presentation([question(1, accept="Толстой", comment="1869")], {"q_on_answer": True})
    prs = Presentation(io.BytesIO(data))
    assert prs.slides[0].notes_slide.notes_text_frame.text == "Ответ: Лев Толстой\nЗачёт: Толстой\nКомментарий: 1869"
    answer_text = [sh.text_frame.text for sh in prs.slides[1].shapes if sh.has_text_frame]
    assert any(t.startswith("Кто написал") and "Ответ: Лев Толстой" in t for t in answer_text)

    without = Presentation(io.BytesIO(g.generate_presentation([question(1)], {"notes": False})))
    assert not without.slides[0].has_notes_slide


def test_notes_master_is_registered_in_presentation_xml():
    """Без <p:notesMasterIdLst> iPhone (Quick Look) показывает белый экран вместо слайдов."""
    data = g.generate_presentation([question(1)], {"notes": True})
    with zipfile.ZipFile(io.BytesIO(data)) as zf:
        prs = etree.fromstring(zf.read("ppt/presentation.xml"))
        rels = etree.fromstring(zf.read("ppt/_rels/presentation.xml.rels"))
    children = [etree.QName(el).localname for el in prs]
    assert children[:3] == ["sldMasterIdLst", "notesMasterIdLst", "sldIdLst"]
    rid = prs.find(f"{{{g.P}}}notesMasterIdLst/{{{g.P}}}notesMasterId").get(f"{{{g.R_NS}}}id")
    target = {r.get("Id"): r.get("Target") for r in rels}[rid]
    assert target.endswith("notesMasters/notesMaster1.xml")

    without = g.generate_presentation([question(1)], {"notes": False})
    with zipfile.ZipFile(io.BytesIO(without)) as zf:
        assert b"notesMasterIdLst" not in zf.read("ppt/presentation.xml")


def test_custom_template_keeps_its_colors_and_drops_its_notes():
    tpl = g.validate_template(custom_template(with_notes=True))
    assert tpl.question.label_id is not None
    data = g.write_presentation(tpl, g.build_specs([question(1)], {"theme": "template", "notes": False}, tpl))
    with zipfile.ZipFile(io.BytesIO(data)) as zf:
        names = zf.namelist()
        slide = etree.fromstring(zf.read("ppt/slides/slide1.xml"))
    assert not [n for n in names if n.startswith("ppt/notesSlides/")]
    assert not slide.findall(f".//{{{g.A}}}srgbClr")  # «как в шаблоне» — цвета не навязываем

    prs = Presentation(io.BytesIO(data))
    texts = [sh.text_frame.text for sh in prs.slides[0].shapes if sh.has_text_frame]
    assert "Вопрос 1" in texts and "Кто написал «Войну и мир»?" in texts


def test_dark_theme_colors_label_and_background():
    data = g.generate_presentation([question(1)], {"theme": "dark"})
    with zipfile.ZipFile(io.BytesIO(data)) as zf:
        slide = zf.read("ppt/slides/slide1.xml").decode()
        rels = zf.read("ppt/slides/_rels/slide1.xml.rels").decode()
    assert '<a:srgbClr val="1E1E1E"/>' in slide and slide.count('val="FFFFFF"') >= 2  # фон, текст и метка
    assert "logo_dark.png" in rels


@pytest.mark.parametrize("data", [b"garbage", custom_template()[:200]])
def test_bad_templates_raise_readable_error(data):
    with pytest.raises(g.TemplateError):
        g.validate_template(data)


# ─── Предпросмотр ────────────────────────────────────────────────────────────

@pytest.mark.skipif(not preview.available(), reason="нет TTF-шрифта")
def test_preview_renders_png_of_slide_proportions():
    from PIL import Image

    tpl = g.standard_template()
    spec = g.build_specs([question(1)], {"theme": "dark"}, tpl)[0]
    image = Image.open(io.BytesIO(preview.render(tpl, spec, dark_logo=True, width=640)))
    assert image.size == (640, 360)
    assert image.getpixel((5, 300)) == (0x1E, 0x1E, 0x1E)


# ─── Таблица и раздатка ──────────────────────────────────────────────────────

def test_scoring_table_has_tour_sums_totals_and_ranks():
    qs = [question(1, tour="Тур 1"), question(2, tour="Тур 1"), question(3, tour="Тур 2")]
    with zipfile.ZipFile(io.BytesIO(extras.scoring_table(qs, "seq", "Кубок"))) as zf:
        sheet = zf.read("xl/worksheets/sheet1.xml").decode()
        strings = zf.read("xl/sharedStrings.xml").decode()
    assert "RANK(" in sheet and sheet.count("SUM(") >= 3 * 20
    assert "Σ Тур 1" in strings and "Итого" in strings


def test_text_handout_detection():
    assert extras.text_handout("Смотрите. [Раздаточный материал: ЖИ-ШИ] Что это?") == "ЖИ-ШИ"
    assert extras.text_handout("Раздаточный материал: РОЗА\nЧто это?") == "РОЗА"
    assert extras.handout_document([question(1)], "seq", "Кубок") is None


# ─── Проверка ────────────────────────────────────────────────────────────────

def test_local_review_finds_repeats_and_suspicious_answers():
    label = lambda q: f"Вопрос {q['number']}"  # noqa: E731
    qs = [
        question(1, "Назовите автора романа «Война и мир», написанного в 1869 году."),
        question(2, "Назовите автора романа «Война и мир», написанного в 1869 году!"),
        question(3, "Что это?", answer="Очень длинный ответ " * 10),
        question(4, "Столица Франции?", answer="Париж"),
    ]
    past = [{"name": "Весенний кубок", "questions": [question(9, "Какая столица у Франции?", answer="париж")]}]
    issues = review.local_issues(qs, label, past)
    text = "\n".join(issues)
    assert "Вопрос 1 и Вопрос 2 почти одинаковые" in text
    assert "подозрительно длинный" in text
    assert "Вопрос 4 уже был в пакете «Весенний кубок»" in text


def test_ai_review_is_skipped_without_key():
    assert review.ai_issues([question(1)], lambda q: "Вопрос 1") is None


# ─── История ─────────────────────────────────────────────────────────────────

def test_history_keeps_last_ten_and_restores(tmp_path):
    history = History(tmp_path)
    ids = [history.save(7, {"name": f"Пакет {i}", "questions": [question(1)]}) for i in range(12)]
    items = history.entries(7)
    assert len(items) == 10 and items[0]["name"] == "Пакет 11"
    assert history.load(7, ids[0]) is None                    # самые старые удалены
    assert history.latest(7)["name"] == "Пакет 11"
    assert len(history.others(7, ids[-1])) == 9
    assert history.load(7, "../../etc") is None               # путь из кнопки не выходит за папку
