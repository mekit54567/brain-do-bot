"""Тесты парсера: python -m pytest"""

import io

import pytest
from docx import Document

import parser as p


@pytest.fixture(autouse=True)
def no_ai(monkeypatch):
    monkeypatch.setattr(p, "GROQ_API_KEY", "")


def make_docx(tmp_path, paragraphs, table=None):
    doc = Document()
    for text in paragraphs:
        doc.add_paragraph(text)
    if table:
        t = doc.add_table(rows=len(table), cols=1)
        for i, text in enumerate(table):
            t.cell(i, 0).text = text
    path = tmp_path / "pack.docx"
    doc.save(path)
    return str(path)


def test_basic_format(tmp_path):
    path = make_docx(tmp_path, [
        "Синхронный турнир 2026",
        "Вопрос 1: Кто написал «Войну и мир»?",
        "Ответ: Лев Толстой",
        "Комментарий: Написал в 1869 году.",
        "Вопрос 2: * Назовите столицу Мозамбика.",
        "Ответ: Мапуту",
    ])
    r = p.parse_file(path)
    assert [q["answer"] for q in r.questions] == ["Лев Толстой", "Мапуту"]
    assert r.questions[0]["comment"] == "Написал в 1869 году."
    assert r.questions[1]["hard"] and r.questions[1]["question"] == "Назовите столицу Мозамбика."
    assert not r.skipped


def test_answer_on_next_line_and_accept(tmp_path):
    path = make_docx(tmp_path, ["Вопрос 1. Столица?", "Ответ:", "Мапуту", "Зачет: Maputo"])
    q = p.parse_file(path).questions[0]
    assert q["answer"] == "Мапуту"
    assert q["accept"] == "Maputo"


def test_year_and_similar_words_do_not_break_question(tmp_path):
    path = make_docx(tmp_path, [
        "3. Какой газ преобладает в атмосфере?",
        "1812 год тут ни при чём.",
        "Ответственность — тоже часть вопроса.",
        "Ответ: Азот",
        "Комментарий: около 78%.",
        "Вторая строка, 12.5 процентов.",
        "Автор: Никита",
    ])
    r = p.parse_file(path)
    assert len(r.questions) == 1
    q = r.questions[0]
    assert "1812 год" in q["question"] and "Ответственность" in q["question"]
    assert q["answer"] == "Азот"
    assert q["comment"].endswith("12.5 процентов.")
    assert q["orig_number"] == 3


def test_tours_tables_and_skipped(tmp_path):
    path = make_docx(
        tmp_path,
        ["Тур 1", "Вопрос 1. Первый?", "Ответ: Да", "2 тур", "Вопрос 1. Второй?", "Ответ: Нет",
         "Вопрос 2. Без ответа"],
        table=["Вопрос 3. Из таблицы?", "Ответ: Таблица\nКомментарий: мягкий перенос"],
    )
    r = p.parse_file(path)
    assert r.tours == ["Тур 1", "Тур 2"]
    assert [q["tour"] for q in r.questions] == ["Тур 1", "Тур 2", "Тур 2"]
    assert r.questions[2]["comment"] == "мягкий перенос"
    assert r.skipped == ["№2 — нет ответа"]
    assert [q["number"] for q in r.questions] == [1, 2, 3]


def test_pictures_are_extracted(tmp_path, png_bytes):
    doc = Document()
    doc.add_paragraph("Вопрос 1. Что на картинке?")
    doc.add_paragraph().add_run().add_picture(io.BytesIO(png_bytes))
    doc.add_paragraph("Ответ: Точка")
    path = tmp_path / "pic.docx"
    doc.save(path)

    q = p.parse_file(str(path)).questions[0]
    assert q["question"] == "Что на картинке?"
    assert len(q["q_pictures"]) == 1 and q["q_pictures"][0].ext == "png"


def test_word_numbered_list_starts_questions_but_bullets_do_not(tmp_path):
    doc = Document()
    doc.add_paragraph("Кто написал «Войну и мир»?", style="List Number")
    doc.add_paragraph("Ответ: Лев Толстой")
    doc.add_paragraph("Комментарий: варианты:")
    doc.add_paragraph("пункт маркированного списка номер один", style="List Bullet")
    doc.add_paragraph("Назовите столицу Мозамбика.", style="List Number")
    doc.add_paragraph("Ответ: Мапуту")
    path = tmp_path / "lists.docx"
    doc.save(path)

    r = p.parse_file(str(path))
    assert [q["answer"] for q in r.questions] == ["Лев Толстой", "Мапуту"]
    assert [q["orig_number"] for q in r.questions] == [1, 2]
    assert "маркированного" in r.questions[0]["comment"]
    assert not r.skipped


def test_plain_text_and_inline_markers():
    text = "1. Столица Франции? Ответ: Париж\n2) Столица Италии?\nОтвет: Рим\nКомментарий: вечный город"
    assert p.looks_like_questions(text)
    r = p.parse_text(text)
    assert [(q["question"], q["answer"]) for q in r.questions] == [("Столица Франции?", "Париж"), ("Столица Италии?", "Рим")]
    assert r.questions[1]["comment"] == "вечный город"


def test_plain_chat_message_is_not_a_pack():
    assert not p.looks_like_questions("Привет! Как дела?")


def test_join_merges_broken_lines():
    assert p._join("Героиня шутки корыстно", "уточняет у кавалера") == "Героиня шутки корыстно уточняет у кавалера"
    assert p._join("Первое предложение.", "Второе") == "Первое предложение.\nВторое"
