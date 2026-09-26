"""Сценарии бота целиком — через подставной Telegram (tests/fakes.py)."""

import asyncio
import io
import zipfile

import pytest
from docx import Document
from pptx import Presentation
from telegram.ext import ApplicationHandlerStop

import access
import admin
import handlers
import keyboards as kb
import preview
import session
from tests.fakes import Harness, buttons, callback_data, make_png


def run(coro):
    return asyncio.run(coro)


def make_pack() -> bytes:
    doc = Document()
    for text in [
        "Тур 1",
        "Вопрос 1: Кто написал «Войну и мир»?",
        "Ответ: Лев Толстой",
        "Комментарий: Написал в 1869 году.",
        "Вопрос 2: * Назовите столицу Мозамбика. [Раздаточный материал: Мапуту — Maputo]",
        "Ответ: Мапуту",
        "Тур 2",
        "Вопрос 1. Что на картинке?",
    ]:
        doc.add_paragraph(text)
    doc.add_paragraph().add_run().add_picture(io.BytesIO(make_png(40, 30)))
    doc.add_paragraph("Ответ: Синий прямоугольник")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def make_template() -> bytes:
    prs = Presentation()  # 4:3, плейсхолдеры без собственных координат
    for body in ("Текст вопроса", "Ответ: пример"):
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = "Вопрос 1"
        slide.placeholders[1].text = body
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


async def press(h: Harness, data: str, message):
    update, answers = h.query_update(data, message)
    await handlers.handle_callback(update, h.ctx())
    assert len(answers) == 1, f"{data}: ответов на кнопку {len(answers)}"
    return answers[0]


async def open_pack(h: Harness):
    await handlers.handle_document(h.document_update("Кубок.docx", make_pack()), h.ctx())
    card = [m for action, m in h.bot.log if action == "edit"][-1]
    assert "Нашёл <b>3</b> вопроса" in card.text
    return card


def pptx_texts(data: bytes) -> list[str]:
    prs = Presentation(io.BytesIO(data))
    return [" / ".join(sh.text_frame.text for sh in s.shapes if sh.has_text_frame) for s in prs.slides]


# ─── Основной сценарий ───────────────────────────────────────────────────────

def test_document_card_generate_and_stats(bot_env):
    h = Harness(bot_env)

    async def scenario():
        card = await open_pack(h)
        assert "🚀 Создать презентацию" in buttons(card.reply_markup)
        assert "📑 Все туры" in buttons(card.reply_markup)

        await press(h, "s:theme", card)
        await press(h, "s:tour", card)          # Тур 1
        assert "📑 Тур 1 (2 вопроса)" in buttons(card.reply_markup)
        assert (await press(h, "gen", card))[0] == "Поехали! 🚀"

        doc = h.bot.last(kind="document")
        assert doc.filename == "Кубок — Тур 1.pptx"
        assert "Готово" in doc.caption and "заметках ведущего" in doc.caption
        texts = pptx_texts(doc.data)
        assert len(texts) == 4 and "Лев Толстой" in texts[1]
        notes = Presentation(io.BytesIO(doc.data)).slides[0].notes_slide.notes_text_frame.text
        assert notes.startswith("Ответ: Лев Толстой")
        assert card.deleted

    run(scenario())
    assert session.stats.user(7) == {"presentations": 1, "questions": 2}
    assert session.history.list(7)[0]["name"] == "Кубок"


def test_busy_and_legacy_buttons(bot_env):
    h = Harness(bot_env)

    async def scenario():
        card = await open_pack(h)
        session.BUSY.add(7)
        assert (await press(h, "gen", card))[0].startswith("Уже делаю")
        session.BUSY.discard(7)
        assert (await press(h, "theme_white", card))[1] is True  # старая кнопка → alert

    run(scenario())


def test_restores_last_pack_after_restart(bot_env):
    h = Harness(bot_env)

    async def scenario():
        card = await open_pack(h)
        session.DOCS.clear()                      # как после перезапуска
        await press(h, "gen", card)
        assert h.bot.last(kind="document").filename == "Кубок.pptx"

    run(scenario())


# ─── Предпросмотр, «Ещё», служебные слайды ───────────────────────────────────

@pytest.mark.skipif(not preview.available(), reason="нет TTF-шрифта для предпросмотра")
def test_preview_album_and_card_below(bot_env):
    h = Harness(bot_env)

    async def scenario():
        card = await open_pack(h)
        await press(h, "pv", card)
        photos = h.bot.sent(kind="photo")
        assert 3 <= len(photos) <= 4 and photos[0].caption.startswith("👁 Предпросмотр")
        assert all(p.photo[:8] == b"\x89PNG\r\n\x1a\n" for p in photos)
        assert card.deleted and "🚀 Создать презентацию" in buttons(h.bot.last(kind="text").reply_markup)

    run(scenario())


def test_more_menu_service_slides_title_scoring_handout_review(bot_env):
    h = Harness(bot_env)

    async def scenario():
        card = await open_pack(h)
        await press(h, "more", card)
        assert "📊 Таблица результатов" in buttons(card.reply_markup)
        assert "🖨 Раздатка" in buttons(card.reply_markup)

        await press(h, "m:service", card)
        await press(h, "m:qa", card)
        assert "title" in callback_data(card.reply_markup)
        await press(h, "title", card)
        await handlers.handle_text(h.text_update("Кубок Брейн-ду 2026"), h.ctx())
        assert "Кубок Брейн-ду 2026" in h.bot.sent(kind="text")[-2].text

        await press(h, "score", card)
        table = h.bot.last(kind="document")
        assert table.filename == "Кубок — таблица.xlsx"
        assert "xl/worksheets/sheet1.xml" in zipfile.ZipFile(io.BytesIO(table.data)).namelist()

        await press(h, "handout", card)
        handout = h.bot.last(kind="document")
        assert handout.filename == "Кубок — раздатка.docx"
        paragraphs = [p.text for p in Document(io.BytesIO(handout.data)).paragraphs if p.text]
        assert "Мапуту — Maputo" in paragraphs

        await press(h, "review", card)
        assert "Проверка «Кубок»" in h.bot.last(kind="text").text

        await press(h, "gen", card)
        texts = pptx_texts(h.bot.last(kind="document").data)
        assert texts[0] == "Кубок Брейн-ду 2026" and texts[1] == "Тур 1" and texts[-1] == "Спасибо за игру!"
        assert "Кто написал" in texts[3]  # вопрос на слайде ответа

    run(scenario())


# ─── Правка вопросов ─────────────────────────────────────────────────────────

def test_edit_mark_and_delete_questions(bot_env):
    h = Harness(bot_env)

    async def scenario():
        card = await open_pack(h)
        await press(h, "list:0", card)
        assert "ed:0" in callback_data(card.reply_markup)

        await press(h, "ed:0", card)
        await press(h, "ef:0:answer", card)
        assert "Сейчас" in h.bot.last(kind="text").text
        await handlers.handle_text(h.text_update("Л. Н. Толстой"), h.ctx())
        assert session.doc_for(7)["questions"][0]["answer"] == "Л. Н. Толстой"

        await press(h, "ef:0:comment", card)
        await handlers.handle_text(h.text_update("-"), h.ctx())
        assert session.doc_for(7)["questions"][0]["comment"] is None

        assert (await press(h, "eh:0", card))[0] == "★ Сложный"
        await press(h, "edel:1", card)
        assert "Точно удалить" in card.text
        await press(h, "edel!:1", card)
        assert len(session.doc_for(7)["questions"]) == 2
        assert (await press(h, "ed:5", card))[1] is True  # такого вопроса уже нет

    run(scenario())
    saved = session.history.latest(7)
    assert saved["questions"][0]["answer"] == "Л. Н. Толстой" and len(saved["questions"]) == 2


# ─── История, шаблон, быстрый режим ──────────────────────────────────────────

def test_history_command_opens_pack(bot_env):
    h = Harness(bot_env)

    async def scenario():
        await open_pack(h)
        session.DOCS.clear()
        await handlers.cmd_history(h.text_update("/history"), h.ctx())
        listing = h.bot.last(kind="text")
        doc_id = session.history.list(7)[0]["id"]
        assert f"h:{doc_id}" in callback_data(listing.reply_markup)
        await press(h, f"h:{doc_id}", listing)
        assert "Кубок" in listing.text and session.DOCS[7]["name"] == "Кубок"

    run(scenario())


def test_custom_template_upload_and_use(bot_env):
    h = Harness(bot_env)

    async def scenario():
        await handlers.handle_document(h.document_update("мой.pptx", make_template()), h.ctx())
        offer = h.bot.last(kind="photo") if preview.available() else h.bot.log[-1][1]
        assert "tpl:ok" in callback_data(offer.reply_markup)
        await press(h, "tpl:ok", offer)
        assert session.has_template(7) and session.settings(h.ctx(), 7)["theme"] == "template"

        card = await open_pack(h)
        await press(h, "gen", card)
        prs = Presentation(io.BytesIO(h.bot.last(kind="document").data))
        assert prs.slide_width * 3 == prs.slide_height * 4   # шаблон 4:3, а не стандартный 16:9

        await handlers.cmd_settings(h.text_update("/settings"), h.ctx())
        settings_msg = h.bot.last(kind="text")
        await press(h, "d:tpl_reset", settings_msg)
        assert not session.has_template(7) and session.settings(h.ctx(), 7)["theme"] == "white"

    run(scenario())


def test_broken_template_is_explained(bot_env):
    h = Harness(bot_env)
    run(handlers.handle_document(h.document_update("битый.pptx", b"not a pptx"), h.ctx()))
    assert "Не получилось открыть" in [m for a, m in h.bot.log if a == "edit"][-1].text


def test_instant_mode_with_pasted_text(bot_env):
    h = Harness(bot_env)

    async def scenario():
        await handlers.cmd_settings(h.text_update("/settings"), h.ctx())
        await press(h, "d:instant", h.bot.last(kind="text"))
        await handlers.handle_text(h.text_update("1. Столица Франции?\nОтвет: Париж"), h.ctx())
        assert h.bot.last(kind="document").filename == "Вопросы из сообщения.pptx"

    run(scenario())


# ─── Доступ и админка ────────────────────────────────────────────────────────

@pytest.fixture
def private_mode(monkeypatch):
    monkeypatch.setattr(access, "ACCESS_MODE", "private")
    monkeypatch.setattr(access, "ADMIN_IDS", {100})
    monkeypatch.setattr(admin, "ADMIN_IDS", {100})


async def guarded(update, h):
    """Прогоняет апдейт через фильтр доступа; True — пропущен дальше."""
    try:
        await admin.access_guard(update, h.ctx())
        return True
    except ApplicationHandlerStop:
        return False


def test_private_mode_request_and_approve(bot_env, private_mode):
    stranger, boss = Harness(bot_env, 55, "Гость"), Harness(bot_env, 100, "Никита")

    async def scenario():
        assert not await guarded(stranger.text_update("/start"), stranger)
        offer = bot_env.last(55)
        assert "закрытый бот" in offer.text

        update, answers = stranger.query_update("acc:req", offer)
        assert not await guarded(update, stranger)
        request = bot_env.last(100)
        assert "просит доступ" in request.text and "acc:ok:55" in callback_data(request.reply_markup)

        assert await guarded(boss.text_update("/users"), boss)  # админа пропускает
        await press(boss, "acc:ok:55", request)
        assert "Доступ к боту открыт" in bot_env.last(55).text
        assert await guarded(stranger.text_update("/start"), stranger)

    run(scenario())


def test_invite_link_is_single_use(bot_env, private_mode):
    boss, guest, other = Harness(bot_env, 100), Harness(bot_env, 56), Harness(bot_env, 57)

    async def scenario():
        await admin.cmd_invite(boss.text_update("/invite"), boss.ctx())
        link = bot_env.last(100).text.split("\n")[1]
        token = link.split("start=inv_")[1]
        assert await guarded(guest.text_update(f"/start inv_{token}"), guest)
        assert not await guarded(other.text_update(f"/start inv_{token}"), other)
        assert await guarded(other.text_update("/id"), other)  # свой ID можно узнать всегда

    run(scenario())


def test_errors_go_to_admin_once(bot_env, private_mode):
    async def scenario():
        error = ValueError("сломалось")
        await admin.report_error(bot_env, error, "пользователь 7, кнопка «gen»")
        await admin.report_error(bot_env, error, "пользователь 7, кнопка «gen»")
        reports = bot_env.sent(100)
        assert len(reports) == 1 and "ValueError" in reports[0].text

    admin._last_error.clear()
    run(scenario())


# ─── Яндекс Диск ─────────────────────────────────────────────────────────────

def test_yadisk_list_search_and_open(bot_env, monkeypatch):
    import yadisk as yd

    files = [{"name": f"Пакет {i:02d} с очень длинным названием для проверки лимита.docx",
              "path": f"disk:/Очень длинная папка/подпапка/Пакет {i:02d}.docx", "modified": ""} for i in range(19)]

    async def list_files(token):
        if token == "expired":
            raise yd.AuthError()
        return files

    async def download_file(token, path):
        return make_pack()

    monkeypatch.setattr(yd, "list_files", list_files)
    monkeypatch.setattr(yd, "download_file", download_file)
    h = Harness(bot_env)

    async def scenario():
        await handlers.handle_text(h.text_update(kb.BTN_YADISK), h.ctx())
        assert "state=" in h.bot.last(kind="text").reply_markup.inline_keyboard[0][0].url

        session.tokens.set(7, "good")
        await handlers.handle_text(h.text_update(kb.BTN_YADISK), h.ctx())
        listing = h.bot.last(kind="text")
        assert max(len(d.encode()) for d in callback_data(listing.reply_markup)) <= 64

        await press(h, "yd:search", listing)
        await handlers.handle_text(h.text_update("Пакет 05"), h.ctx())
        found = h.bot.last(kind="text")
        assert found.text.startswith("🔎 Поиск «Пакет 05»: 1")
        await press(h, "yd:f:0", found)
        assert "Нашёл <b>3</b> вопроса" in found.text

        session.tokens.set(7, "expired")
        await press(h, "yd:refresh", listing)
        assert "истёк" in listing.text and session.tokens.get(7) is None

    run(scenario())


def test_yadisk_oauth_callback_page(bot_env, monkeypatch):
    import socket

    import httpx

    import bot
    import yadisk as yd

    with socket.socket() as s:
        s.bind(("127.0.0.1", 0))
        port = s.getsockname()[1]
    notified = []

    async def notify(app, user_id, chat_id):
        notified.append((user_id, chat_id))

    monkeypatch.setattr(bot, "SERVER_PORT", port)
    monkeypatch.setattr(handlers, "notify_yadisk_connected", notify)
    monkeypatch.setattr(yd, "exchange_code_for_token", lambda code: f"TOKEN-{code}")

    async def scenario():
        bot.start_web_server(type("App", (), {"bot": bot_env})(), asyncio.get_running_loop())
        base = f"http://127.0.0.1:{port}"
        async with httpx.AsyncClient() as client:
            assert (await client.get(f"{base}/")).text == "OK"
            forged = await client.get(f"{base}/yadisk/callback?code=abc&state=7")
            assert forged.status_code == 400 and "charset=utf-8" in forged.headers["content-type"]

            state = yd.auth_url(7, 42).split("state=")[1].split("&")[0]
            ok = await client.get(f"{base}/yadisk/callback?code=abc&state={state}")
            await asyncio.sleep(0.1)
            assert ok.status_code == 200 and "t.me/brain_do_test_bot" in ok.text
            replay = await client.get(f"{base}/yadisk/callback?code=abc&state={state}")
            assert replay.status_code == 400

    run(scenario())
    assert session.tokens.get(7) == "TOKEN-abc" and notified == [(7, 42)]
