"""
Файлы для игры помимо презентации:
- таблица результатов (.xlsx) — команды × вопросы, суммы по турам, места;
- раздаточный материал (.docx) — картинки и тексты раздаток для печати.
"""

import io
import re

import xlsxwriter
from xlsxwriter.utility import xl_rowcol_to_cell
from docx import Document
from docx.enum.text import WD_BREAK
from docx.shared import Cm, Pt

from generator import question_label

TEAMS = 20

RE_HANDOUT_BRACKETS = re.compile(r"\[\s*Раздаточный материал\s*:?\s*(.+?)\]", re.I | re.S)
RE_HANDOUT_LINE = re.compile(r"^\s*Раздаточный материал\s*:?\s*(.+)$", re.I | re.M)


def number_of(q: dict, numbering: str) -> str:
    label = question_label(q, "seq" if numbering == "none" else numbering)
    return label.removeprefix("Вопрос ").strip()


# ─── Таблица результатов ─────────────────────────────────────────────────────

def scoring_table(questions: list[dict], numbering: str, title: str) -> bytes:
    buf = io.BytesIO()
    wb = xlsxwriter.Workbook(buf, {"in_memory": True})
    ws = wb.add_worksheet("Результаты")

    f_title = wb.add_format({"bold": True, "font_size": 16})
    f_head = wb.add_format({"bold": True, "align": "center", "valign": "vcenter", "border": 1, "bg_color": "#F2F2F2"})
    f_tour = wb.add_format({"bold": True, "align": "center", "border": 1, "bg_color": "#FDE9E7", "font_color": "#B3261E"})
    f_cell = wb.add_format({"align": "center", "border": 1})
    f_team = wb.add_format({"border": 1})
    f_sum = wb.add_format({"bold": True, "align": "center", "border": 1, "bg_color": "#F2F2F2"})
    f_total = wb.add_format({"bold": True, "align": "center", "border": 2, "bg_color": "#FFF2CC"})
    f_plus = wb.add_format({"bg_color": "#C6EFCE", "font_color": "#006100"})

    ws.write(0, 0, title, f_title)
    ws.write(1, 0, "Ставь 1 за взятый вопрос — суммы и места посчитаются сами.")

    tours = []  # [(название, первая колонка, последняя колонка)]
    head_row, first_row = 3, 4
    last_row = first_row + TEAMS - 1
    col = 2
    for q in questions:
        tour = q.get("tour") or ""
        if not tours or tours[-1][0] != tour:
            tours.append([tour, col, col])
        tours[-1][2] = col
        ws.write(head_row, col, number_of(q, numbering), f_head)
        col += 1
    q_last = col - 1

    named_tours = [t for t in tours if t[0]]
    if len(named_tours) > 1:
        for name, c1, c2 in named_tours:
            if c1 == c2:
                ws.write(head_row - 1, c1, name, f_tour)
            else:
                ws.merge_range(head_row - 1, c1, head_row - 1, c2, name, f_tour)

    sum_cols = []
    if len(named_tours) > 1:
        for name, c1, c2 in named_tours:
            ws.write(head_row, col, f"Σ {name}", f_head)
            sum_cols.append((col, c1, c2))
            col += 1
    total_col = col

    ws.write(head_row, 0, "Место", f_head)
    ws.write(head_row, 1, "Команда", f_head)
    ws.write(head_row, total_col, "Итого", f_head)

    for r in range(first_row, last_row + 1):
        ws.write(r, 1, "", f_team)
        for c in range(2, q_last + 1):
            ws.write_blank(r, c, None, f_cell)
        for c, c1, c2 in sum_cols:
            ws.write_formula(r, c, f"=SUM({_cell(r, c1)}:{_cell(r, c2)})", f_sum, 0)
        ws.write_formula(r, total_col, f"=SUM({_cell(r, 2)}:{_cell(r, q_last)})", f_total, 0)
        total_range = f"{_cell(first_row, total_col, abs_=True)}:{_cell(last_row, total_col, abs_=True)}"
        ws.write_formula(
            r, 0, f'=IF({_cell(r, 1)}="","",RANK({_cell(r, total_col)},{total_range}))', f_cell, ""
        )

    ws.conditional_format(first_row, 2, last_row, q_last, {"type": "cell", "criteria": ">", "value": 0, "format": f_plus})
    ws.set_column(0, 0, 7)
    ws.set_column(1, 1, 26)
    ws.set_column(2, q_last, 4.5)
    ws.set_column(q_last + 1, total_col, 10)
    ws.freeze_panes(first_row, 2)
    ws.set_landscape()
    ws.fit_to_pages(1, 1)

    answers = wb.add_worksheet("Ответы")
    f_wrap = wb.add_format({"text_wrap": True, "valign": "top", "border": 1})
    for c, name in enumerate(["№", "Вопрос", "Ответ", "Зачёт"]):
        answers.write(0, c, name, f_head)
    for r, q in enumerate(questions, 1):
        text = q["question"].replace("\n", " ")
        answers.write(r, 0, number_of(q, numbering), f_cell)
        answers.write(r, 1, text[:150] + ("…" if len(text) > 150 else ""), f_wrap)
        answers.write(r, 2, q["answer"], f_wrap)
        answers.write(r, 3, q.get("accept") or "", f_wrap)
    answers.set_column(0, 0, 6)
    answers.set_column(1, 1, 70)
    answers.set_column(2, 3, 28)
    answers.freeze_panes(1, 0)

    wb.close()
    return buf.getvalue()


def _cell(row: int, col: int, abs_: bool = False) -> str:
    return xl_rowcol_to_cell(row, col, row_abs=abs_, col_abs=abs_)


# ─── Раздаточный материал ────────────────────────────────────────────────────

def text_handout(question: str) -> str | None:
    """Текст раздатки: «[Раздаточный материал: …]» или строка «Раздаточный материал: …»."""
    m = RE_HANDOUT_BRACKETS.search(question) or RE_HANDOUT_LINE.search(question)
    return m.group(1).strip() if m else None


def has_handouts(questions: list[dict]) -> bool:
    return any(q.get("q_pictures") or text_handout(q["question"]) for q in questions)


def handout_document(questions: list[dict], numbering: str, title: str) -> bytes | None:
    """Раздатка для печати: каждый вопрос с картинкой или текстом раздатки — на своей странице."""
    items = [q for q in questions if q.get("q_pictures") or text_handout(q["question"])]
    if not items:
        return None

    doc = Document()
    for section in doc.sections:
        section.left_margin = section.right_margin = Cm(1.5)
        section.top_margin = section.bottom_margin = Cm(1.5)
    usable_w, usable_h = Cm(18), Cm(23)

    for i, q in enumerate(items):
        head = doc.add_paragraph()
        run = head.add_run(f"{title} · Вопрос {number_of(q, numbering)}")
        run.bold = True
        run.font.size = Pt(14)

        handout = text_handout(q["question"])
        if handout:
            p = doc.add_paragraph()
            p.add_run(handout).font.size = Pt(28)

        pictures = q.get("q_pictures") or []
        for pic in pictures:
            ratio = pic.height / max(pic.width, 1)
            width = usable_w
            max_h = usable_h // max(len(pictures), 1)
            if width * ratio > max_h:
                width = int(max_h / ratio)
            doc.add_paragraph().add_run().add_picture(io.BytesIO(pic.blob), width=width)

        if i < len(items) - 1:
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()
