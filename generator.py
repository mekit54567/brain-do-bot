"""
Генератор презентаций Brain-Do.

1. Шаблон (стандартный template.pptx или свой .pptx пользователя) разбирается
   один раз: 1-й слайд — образец вопроса, 2-й — образец ответа. На каждом
   ищем большое текстовое поле (тело) и маленькое с «Вопрос N» (метку).
2. build_specs() раскладывает вопросы по слайдам: какой текст, каким кеглем,
   где картинки, что в заметках. Этим же описанием пользуется предпросмотр.
3. write_presentation() клонирует слайды шаблона в памяти и собирает .pptx
   без внешних скриптов и «мёртвых» слайдов шаблона внутри файла.
"""

import hashlib
import io
import math
import re
import zipfile
from dataclasses import dataclass, field
from functools import lru_cache

from lxml import etree

from config import BASE_DIR

TEMPLATE_PATH = BASE_DIR / "template.pptx"
LOGO_DARK_PATH = BASE_DIR / "assets" / "logo_dark.png"

A = "http://schemas.openxmlformats.org/drawingml/2006/main"
P = "http://schemas.openxmlformats.org/presentationml/2006/main"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PKG_REL = "http://schemas.openxmlformats.org/package/2006/relationships"
CT_NS = "http://schemas.openxmlformats.org/package/2006/content-types"

SLIDE_CT = "application/vnd.openxmlformats-officedocument.presentationml.slide+xml"
IMAGE_CT = {"png": "image/png", "jpg": "image/jpeg", "jpeg": "image/jpeg", "gif": "image/gif", "bmp": "image/bmp"}
RE_SLIDE_PART = re.compile(r"^ppt/slides/(_rels/)?slide\d+\.xml(\.rels)?$")
RE_DROPPED_PART = re.compile(r"^ppt/(notesSlides|comments)/")  # части, привязанные к слайдам шаблона
DROPPED_REL_TYPES = ("/notesSlide", "/comments", "/tags")
RE_LABEL = re.compile(r"^\s*Вопрос\s*№?\s*\d*\s*$", re.I)

EMU_PER_PT = 12700
CHAR_WIDTH = 0.5              # средняя ширина символа Calibri в долях кегля (замер в PowerPoint)
LINE_HEIGHT = 1.1             # высота строки при интервале 90% (замер: 1.08) + запас
FRAME_PADDING_PT = 12         # внутренние поля рамки по вертикали
PICTURE_GAP = 250_000         # отступ между текстом и картинкой, EMU
MIN_PT = 16

THEMES = {
    # «Как в шаблоне» — только для своего шаблона: цвета и фон не трогаем
    "template": {"label": "🖼 Как в шаблоне", "bg": None, "text": None, "muted": None, "accent": "CC0000"},
    "white": {"label": "⬜ Белая", "bg": None, "text": "000000", "muted": "7F7F7F", "accent": "CC0000"},
    "dark": {"label": "⬛ Тёмная", "bg": "1E1E1E", "text": "FFFFFF", "muted": "A6A6A6", "accent": "FF5A4E"},
}


def qn(ns: str, tag: str) -> str:
    return f"{{{ns}}}{tag}"


class TemplateError(Exception):
    """Шаблон не подходит — текст понятен пользователю."""


# ─── Описание слайдов ────────────────────────────────────────────────────────

@dataclass
class Box:
    x: int
    y: int
    w: int
    h: int


@dataclass
class Run:
    text: str
    bold: bool = False
    color: str | None = None     # hex; None — цвет из шаблона


@dataclass
class Para:
    runs: list[Run]
    size: float                  # кегль, pt
    align: str = "l"             # l | just | ctr


@dataclass
class SlideSpec:
    kind: str                    # question | answer | title | tour | final
    source: str                  # какой слайд шаблона клонировать: question | answer
    label: str | None            # «Вопрос N»; None — метку убрать
    body: Box
    paras: list[Para]
    anchor: str = "t"            # t | ctr
    pictures: list[tuple] = field(default_factory=list)  # [(Picture, Box)]
    notes: str | None = None
    bg: str | None = None        # цвет фона темы
    text_color: str | None = None  # цвет метки «Вопрос N»


# ─── Шаблон ──────────────────────────────────────────────────────────────────

@dataclass
class TemplateSlide:
    path: str                    # ppt/slides/slideN.xml
    xml: bytes
    rels: bytes
    body_id: str
    label_id: str | None
    body: Box
    label: Box | None


@dataclass
class Template:
    files: dict
    width: int
    height: int
    question: TemplateSlide
    answer: TemplateSlide
    standard: bool
    digest: str


@lru_cache(maxsize=1)
def standard_template() -> Template:
    return _parse_template(TEMPLATE_PATH.read_bytes(), standard=True)


_custom_cache: dict[str, Template] = {}


def load_template(data: bytes | None) -> Template:
    """Свой шаблон пользователя (с кэшем) или стандартный."""
    if not data:
        return standard_template()
    digest = hashlib.sha1(data).hexdigest()
    if digest not in _custom_cache:
        if len(_custom_cache) > 20:
            _custom_cache.clear()
        _custom_cache[digest] = _parse_template(data, standard=False)
    return _custom_cache[digest]


def validate_template(data: bytes) -> Template:
    """Проверяет присланный .pptx и возвращает разобранный шаблон."""
    try:
        return _parse_template(data, standard=False)
    except TemplateError:
        raise
    except Exception as e:
        raise TemplateError("Не получилось открыть файл как презентацию PowerPoint.") from e


def _parse_template(data: bytes, standard: bool) -> Template:
    with zipfile.ZipFile(io.BytesIO(data)) as zf:
        files = {name: zf.read(name) for name in zf.namelist()}
    prs = etree.fromstring(files["ppt/presentation.xml"])
    size = prs.find(qn(P, "sldSz"))
    width, height = int(size.get("cx")), int(size.get("cy"))

    targets = _rels_map(files, "ppt/presentation.xml")
    slide_ids = prs.find(qn(P, "sldIdLst"))
    paths = [_resolve("ppt/presentation.xml", targets[s.get(qn(R_NS, "id"))][1])
             for s in (slide_ids if slide_ids is not None else [])]
    if not paths:
        raise TemplateError("В шаблоне нет ни одного слайда.")

    question = _parse_template_slide(files, paths[0], width, height)
    answer = _parse_template_slide(files, paths[1], width, height) if len(paths) > 1 else question
    return Template(files, width, height, question, answer, standard, hashlib.sha1(data).hexdigest())


def _parse_template_slide(files: dict, path: str, width: int, height: int) -> TemplateSlide:
    root = etree.fromstring(files[path])
    rels_path = _rels_path(path)
    shapes = []
    for sp in root.iter(qn(P, "sp")):
        if sp.find(qn(P, "txBody")) is None:
            continue
        box = _shape_box(files, path, sp) or Box(int(width * 0.05), int(height * 0.15), int(width * 0.9), int(height * 0.75))
        text = "".join(t.text or "" for t in sp.iter(qn(A, "t")))
        shapes.append((sp, box, text))
    if not shapes:
        raise TemplateError("На слайдах шаблона нет текстового поля для вопроса.")

    labels = [s for s in shapes if RE_LABEL.match(s[2])]
    label = min(labels, key=lambda s: s[1].w * s[1].h) if labels else None
    candidates = [s for s in shapes if s is not label] or shapes
    body = max(candidates, key=lambda s: s[1].w * s[1].h)
    if label is not None and label[1].w * label[1].h >= body[1].w * body[1].h:
        label = None

    return TemplateSlide(
        path=path,
        xml=files[path],
        rels=files.get(rels_path, _empty_rels()),
        body_id=_shape_id(body[0]),
        label_id=_shape_id(label[0]) if label else None,
        body=body[1],
        label=label[1] if label else None,
    )


def placeholder_chain(files: dict, part: str, sp) -> list[tuple[str, object]]:
    """Плейсхолдеры макета и мастера, от которых фигура наследует положение и стиль."""
    ph = sp.find(f"{qn(P, 'nvSpPr')}/{qn(P, 'nvPr')}/{qn(P, 'ph')}")
    if ph is None:
        return []
    chain = []
    current = part
    for rel_type in ("/slideLayout", "/slideMaster"):
        parent = next((t for typ, t in _rels_map(files, current).values() if typ.endswith(rel_type)), None)
        if not parent:
            break
        current = _resolve(current, parent)
        for candidate in etree.fromstring(files[current]).iter(qn(P, "sp")):
            c_ph = candidate.find(f"{qn(P, 'nvSpPr')}/{qn(P, 'nvPr')}/{qn(P, 'ph')}")
            if c_ph is None:
                continue
            same_idx = ph.get("idx") and c_ph.get("idx") == ph.get("idx")
            if same_idx or _ph_kind(c_ph.get("type")) == _ph_kind(ph.get("type")):
                chain.append((current, candidate))
                break
    return chain


def master_path(files: dict, slide_path: str) -> str | None:
    """Мастер-слайд, к которому через макет привязан слайд."""
    current = slide_path
    for rel_type in ("/slideLayout", "/slideMaster"):
        parent = next((t for typ, t in _rels_map(files, current).values() if typ.endswith(rel_type)), None)
        if not parent:
            return None
        current = _resolve(current, parent)
    return current


def _shape_box(files: dict, part: str, sp) -> Box | None:
    """Положение фигуры; у плейсхолдеров без своего xfrm — из макета или мастера."""
    for el in [sp] + [candidate for _, candidate in placeholder_chain(files, part, sp)]:
        xfrm = el.find(f"{qn(P, 'spPr')}/{qn(A, 'xfrm')}")
        if xfrm is not None and xfrm.find(qn(A, "off")) is not None:
            off, ext = xfrm.find(qn(A, "off")), xfrm.find(qn(A, "ext"))
            return Box(int(off.get("x")), int(off.get("y")), int(ext.get("cx")), int(ext.get("cy")))
    return None


def _ph_kind(ph_type: str | None) -> str:
    """Мастер знает только title и body — сводим к ним подтипы."""
    if ph_type in ("title", "ctrTitle"):
        return "title"
    if ph_type in (None, "body", "subTitle", "obj"):
        return "body"
    return ph_type


# ─── Раскладка вопросов по слайдам ───────────────────────────────────────────

def question_label(q: dict, numbering: str) -> str:
    if numbering == "none":
        return "Вопрос"
    if numbering == "orig":
        return f"Вопрос {q.get('orig_number') or q['number']}"
    return f"Вопрос {q['number']}"


def build_specs(questions: list[dict], settings: dict, template: Template, title: str | None = None) -> list[SlideSpec]:
    theme = THEMES.get(settings.get("theme"), THEMES["white"])
    numbering = settings.get("numbering", "seq")
    service = settings.get("service", False)
    specs: list[SlideSpec] = []

    if service:
        specs.append(_service_spec("title", title or "Brain-Do", template, theme, max_pt=54))

    previous_tour = object()
    for q in questions:
        tour = q.get("tour")
        if service and not settings.get("shuffle") and tour and tour != previous_tour:
            specs.append(_service_spec("tour", tour, template, theme, max_pt=66))
        previous_tour = tour

        label = question_label(q, numbering)
        specs.append(_question_spec(q, label, settings, template, theme))
        specs.append(_answer_spec(q, label, settings, template, theme))

    if service:
        specs.append(_service_spec("final", "Спасибо за игру!", template, theme, max_pt=60))
    return specs


def _question_spec(q: dict, label: str, settings: dict, template: Template, theme: dict) -> SlideSpec:
    slide = template.question
    pictures = q.get("q_pictures") or []
    text_box, placed = _layout_pictures(slide.body, pictures, text_too=bool(q["question"]))
    align = "l" if pictures else "just"

    blocks = [(line, 1.0, False, "text") for line in q["question"].split("\n") if line]
    if q.get("hard"):
        blocks.append(("★ Сложный вопрос", 0.5, True, "accent"))
    paras = _make_paras(blocks, text_box, theme, max_pt=44, align=align)

    notes = _answer_notes(q) if settings.get("notes", True) else None
    return SlideSpec("question", "question", label, text_box, paras, pictures=placed, notes=notes,
                     bg=theme["bg"], text_color=theme["text"])


def _answer_spec(q: dict, label: str, settings: dict, template: Template, theme: dict) -> SlideSpec:
    slide = template.answer
    pictures = q.get("a_pictures") or []
    text_box, placed = _layout_pictures(slide.body, pictures, text_too=True)

    blocks = []
    if settings.get("q_on_answer") and q["question"]:
        text = q["question"].replace("\n", " ")
        text = text if len(text) <= 300 else text[:297].rstrip() + "…"
        blocks += [(text, 0.5, False, "muted"), ("", 0.25, False, "text")]
    blocks.append(([("Ответ: ", True), (q["answer"], False)], 1.0, None, "text"))
    if q.get("accept"):
        blocks.append(([("Зачёт: ", True), (q["accept"], False)], 0.8, None, "text"))
    if q.get("comment"):
        lines = q["comment"].split("\n")
        blocks.append(("", 0.5, False, "text"))
        blocks.append(([("Комментарий: ", True), (lines[0], False)], 0.72, None, "text"))
        blocks += [(line, 0.72, False, "text") for line in lines[1:] if line]
    paras = _make_paras(blocks, text_box, theme, max_pt=40, align="l")
    return SlideSpec("answer", "answer", label, text_box, paras, pictures=placed,
                     bg=theme["bg"], text_color=theme["text"])


def _service_spec(kind: str, text: str, template: Template, theme: dict, max_pt: int) -> SlideSpec:
    box = template.question.body
    paras = _make_paras([(text, 1.0, True, "text")], box, theme, max_pt=max_pt, align="ctr")
    return SlideSpec(kind, "question", None, box, paras, anchor="ctr", bg=theme["bg"], text_color=theme["text"])


def _answer_notes(q: dict) -> str:
    lines = [f"Ответ: {q['answer']}"]
    if q.get("accept"):
        lines.append(f"Зачёт: {q['accept']}")
    if q.get("comment"):
        lines.append(f"Комментарий: {q['comment']}")
    return "\n".join(lines)


def _make_paras(blocks: list[tuple], box: Box, theme: dict, max_pt: int, align: str) -> list[Para]:
    """blocks: (текст | [(текст, жирный)], относительный кегль, жирный, цвет-роль)."""
    normalized = []
    for content, rel, bold, role in blocks:
        parts = content if isinstance(content, list) else [(content, bold)]
        normalized.append((parts, rel, role))

    size = _fit_size([("".join(t for t, _ in parts), rel) for parts, rel, _ in normalized], box, max_pt, MIN_PT)
    paras = []
    for parts, rel, role in normalized:
        pt = size * rel
        if role == "accent":
            pt = max(pt, 18)  # пометка «сложный» всегда читаемая
        color = theme["accent"] if role == "accent" else theme["muted"] if role == "muted" else theme["text"]
        paras.append(Para([Run(text, bold, color) for text, bold in parts if text or len(parts) == 1],
                          round(pt), align))
    return paras


def _fit_size(texts: list[tuple[str, float]], box: Box, max_pt: int, min_pt: int) -> int:
    """Самый крупный кегль, при котором текст помещается в рамку."""
    width_pt = (box.w - 2 * 91425) / EMU_PER_PT
    height_pt = box.h / EMU_PER_PT - FRAME_PADDING_PT
    for size in range(max_pt, min_pt - 1, -1):
        total = 0.0
        for text, rel in texts:
            sz = size * rel
            chars_per_line = max(1, int(width_pt / (sz * CHAR_WIDTH)))
            total += _wrapped_lines(text, chars_per_line) * sz * LINE_HEIGHT
        if total <= height_pt:
            return size
    return min_pt


def _wrapped_lines(text: str, chars_per_line: int) -> int:
    """Сколько строк займёт текст при переносе по словам."""
    if not text:
        return 1
    lines, current = 1, 0
    for word in text.split():
        length = len(word)
        if current and current + 1 + length > chars_per_line:
            lines += 1
            current = 0
        if length > chars_per_line:
            lines += math.ceil(length / chars_per_line) - 1
            current = length % chars_per_line
        else:
            current += (1 if current else 0) + length
    return lines


def _layout_pictures(body: Box, pictures: list, text_too: bool) -> tuple[Box, list[tuple]]:
    """Картинки справа от текста (или на всю область). Возвращает (рамка текста, [(картинка, рамка)])."""
    if not pictures:
        return body, []
    if text_too:
        text_w = int(body.w * 0.56)
        area = Box(body.x + text_w + PICTURE_GAP, body.y, body.w - text_w - PICTURE_GAP, body.h)
        text_box = Box(body.x, body.y, text_w, body.h)
    else:
        # Вопрос-картинка: картинка на всю область, под ней узкая полоса для пометок
        pic_h = int(body.h * 0.86)
        area = Box(body.x, body.y, body.w, pic_h)
        text_box = Box(body.x, body.y + pic_h, body.w, body.h - pic_h)

    cell_h = (area.h - PICTURE_GAP * (len(pictures) - 1)) // len(pictures)
    placed = []
    for i, pic in enumerate(pictures):
        scale = min(area.w / max(pic.width, 1), cell_h / max(pic.height, 1))
        cx, cy = int(pic.width * scale), int(pic.height * scale)
        cell_y = area.y + i * (cell_h + PICTURE_GAP)
        placed.append((pic, Box(area.x + (area.w - cx) // 2, cell_y + (cell_h - cy) // 2, cx, cy)))
    return text_box, placed


# ─── Сборка .pptx ────────────────────────────────────────────────────────────

def generate_presentation(questions: list[dict], settings: dict, template: Template | None = None,
                          title: str | None = None) -> bytes:
    """Собирает .pptx и возвращает его содержимое."""
    template = template or standard_template()
    specs = build_specs(questions, settings, template, title)
    return write_presentation(template, specs, dark_logo=settings.get("theme") == "dark")


def write_presentation(template: Template, specs: list[SlideSpec], dark_logo: bool = False) -> bytes:
    files = {name: data for name, data in template.files.items()
             if not RE_SLIDE_PART.match(name) and not RE_DROPPED_PART.match(name)}
    media = _Media(files)
    swap_logo = dark_logo and template.standard
    if swap_logo:
        files["ppt/media/logo_dark.png"] = LOGO_DARK_PATH.read_bytes()

    for i, spec in enumerate(specs, 1):
        slide_xml, rels_xml = _write_slide(template, spec, media, swap_logo)
        files[f"ppt/slides/slide{i}.xml"] = slide_xml
        files[f"ppt/slides/_rels/slide{i}.xml.rels"] = rels_xml

    _update_presentation(files, len(specs))
    _update_content_types(files, len(specs), media.extensions)
    _update_app_props(files, len(specs))

    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("[Content_Types].xml", files.pop("[Content_Types].xml"))
        for name, data in files.items():
            zf.writestr(name, data)
    data = buf.getvalue()

    if any(spec.notes for spec in specs):
        data = _add_notes(data, [spec.notes for spec in specs])
    return data


def _add_notes(data: bytes, notes: list[str | None]) -> bytes:
    """Заметки докладчика: ведущий видит ответ в режиме докладчика, зал — нет."""
    from pptx import Presentation  # тяжёлый импорт — только когда нужен

    prs = Presentation(io.BytesIO(data))
    for slide, text in zip(prs.slides, notes):
        if text:
            slide.notes_slide.notes_text_frame.text = text
    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()


class _Media:
    """Картинки из вопросов: складывает в ppt/media без дублей."""

    def __init__(self, files: dict) -> None:
        self.files = files
        self.by_hash: dict[str, str] = {}
        self.extensions: set[str] = set()

    def add(self, blob: bytes, ext: str) -> str:
        ext = "jpg" if ext == "jpeg" else ext
        digest = hashlib.sha1(blob).hexdigest()[:16]
        if digest not in self.by_hash:
            name = f"bd_img_{len(self.by_hash) + 1}.{ext}"
            self.files[f"ppt/media/{name}"] = blob
            self.by_hash[digest] = name
            self.extensions.add(ext)
        return self.by_hash[digest]


def _write_slide(template: Template, spec: SlideSpec, media: _Media, swap_logo: bool) -> tuple[bytes, bytes]:
    source = template.question if spec.source == "question" else template.answer
    root = etree.fromstring(source.xml)
    rels = etree.fromstring(source.rels)
    for rel in list(rels):
        if rel.get("Type", "").endswith(DROPPED_REL_TYPES):
            rels.remove(rel)

    shapes = {_shape_id(sp): sp for sp in root.iter(qn(P, "sp"))}
    body = shapes[source.body_id]
    _set_box(body, spec.body)
    _fill_body(body.find(qn(P, "txBody")), spec)

    if source.label_id:
        label = shapes[source.label_id]
        if spec.label is None:
            label.getparent().remove(label)
        else:
            _set_label(label.find(qn(P, "txBody")), spec.label, spec.text_color)

    sp_tree = root.find(f"{qn(P, 'cSld')}/{qn(P, 'spTree')}")
    used = [int(r.get("Id")[3:]) for r in rels if r.get("Id", "")[3:].isdigit()]
    next_rid = max(used, default=0) + 1
    for i, (pic, box) in enumerate(spec.pictures):
        name = media.add(pic.blob, pic.ext)
        rid = f"rId{next_rid + i}"
        rel = etree.SubElement(rels, qn(PKG_REL, "Relationship"))
        rel.set("Id", rid)
        rel.set("Type", f"{R_NS}/image")
        rel.set("Target", f"../media/{name}")
        sp_tree.append(_picture_element(5000 + i, rid, box))

    if swap_logo:
        for rel in rels:
            if rel.get("Target", "").endswith("media/image1.png"):
                rel.set("Target", "../media/logo_dark.png")
    if spec.bg:
        _set_background(root, spec.bg)
    return _xml(root), _xml(rels)


def _fill_body(tf, spec: SlideSpec) -> None:
    _clear(tf)
    for para in spec.paras:
        p = _new_paragraph(tf, algn=para.align)
        for run in para.runs:
            if run.text:
                p.append(_make_run(run.text, bold=run.bold, sz=int(para.size * 100), color=run.color))
        # Кегль конца абзаца задаёт высоту пустых строк-отбивок (иначе берётся из шаблона)
        etree.SubElement(p, qn(A, "endParaRPr")).set("sz", str(int(para.size * 100)))
    if tf.find(qn(A, "p")) is None:
        _new_paragraph(tf)

    # Убираем отступы из lstStyle шаблона и включаем «сжать при переполнении»
    for lvl in tf.findall(f"{qn(A, 'lstStyle')}//{qn(A, 'pPr')}"):
        lvl.set("marL", "0")
        lvl.set("indent", "0")
    body_pr = tf.find(qn(A, "bodyPr"))
    if body_pr is not None:
        body_pr.set("anchor", spec.anchor)
        for child in list(body_pr):
            if child.tag.endswith("Autofit"):
                body_pr.remove(child)
        etree.SubElement(body_pr, qn(A, "normAutofit"))


def _set_label(tf, text: str, color: str | None) -> None:
    """Меняем текст метки, сохраняя оформление шаблона (шрифт, кегль, выравнивание)."""
    paragraphs = tf.findall(qn(A, "p"))
    first = paragraphs[0]
    for p in paragraphs[1:]:
        tf.remove(p)
    runs = first.findall(qn(A, "r"))
    if runs:
        run = runs[0]
        for extra in runs[1:]:
            first.remove(extra)
        for el in first:
            if el.tag in (qn(A, "fld"), qn(A, "br")):
                first.remove(el)
        run.find(qn(A, "t")).text = text
        if color:
            r_pr = run.find(qn(A, "rPr"))
            if r_pr is None:
                r_pr = etree.Element(qn(A, "rPr"))
                run.insert(0, r_pr)
            for fill in r_pr.findall(qn(A, "solidFill")):
                r_pr.remove(fill)
            fill = etree.Element(qn(A, "solidFill"))
            etree.SubElement(fill, qn(A, "srgbClr")).set("val", color)
            r_pr.insert(_fill_position(r_pr), fill)
    else:
        first.append(_make_run(text, bold=True, sz=2400, color=color))


def _fill_position(r_pr) -> int:
    """solidFill в rPr должен идти после ln, но до шрифтов (latin, ea, cs…)."""
    for i, child in enumerate(r_pr):
        if etree.QName(child).localname not in ("ln",):
            return i
    return len(r_pr)


def _set_box(sp, box: Box) -> None:
    sp_pr = sp.find(qn(P, "spPr"))
    xfrm = sp_pr.find(qn(A, "xfrm"))
    if xfrm is None:
        xfrm = etree.Element(qn(A, "xfrm"))
        sp_pr.insert(0, xfrm)
    for child in list(xfrm):
        xfrm.remove(child)
    off = etree.SubElement(xfrm, qn(A, "off"))
    off.set("x", str(box.x))
    off.set("y", str(box.y))
    ext = etree.SubElement(xfrm, qn(A, "ext"))
    ext.set("cx", str(box.w))
    ext.set("cy", str(box.h))


def _picture_element(shape_id: int, rid: str, box: Box):
    pic = etree.Element(qn(P, "pic"))
    nv = etree.SubElement(pic, qn(P, "nvPicPr"))
    c_nv = etree.SubElement(nv, qn(P, "cNvPr"))
    c_nv.set("id", str(shape_id))
    c_nv.set("name", f"Picture {shape_id}")
    c_nv_pic = etree.SubElement(nv, qn(P, "cNvPicPr"))
    etree.SubElement(c_nv_pic, qn(A, "picLocks")).set("noChangeAspect", "1")
    etree.SubElement(nv, qn(P, "nvPr"))

    fill = etree.SubElement(pic, qn(P, "blipFill"))
    etree.SubElement(fill, qn(A, "blip")).set(qn(R_NS, "embed"), rid)
    etree.SubElement(etree.SubElement(fill, qn(A, "stretch")), qn(A, "fillRect"))

    sp_pr = etree.SubElement(pic, qn(P, "spPr"))
    xfrm = etree.SubElement(sp_pr, qn(A, "xfrm"))
    off = etree.SubElement(xfrm, qn(A, "off"))
    off.set("x", str(box.x))
    off.set("y", str(box.y))
    ext = etree.SubElement(xfrm, qn(A, "ext"))
    ext.set("cx", str(box.w))
    ext.set("cy", str(box.h))
    geom = etree.SubElement(sp_pr, qn(A, "prstGeom"))
    geom.set("prst", "rect")
    etree.SubElement(geom, qn(A, "avLst"))
    return pic


def _set_background(root, color: str) -> None:
    c_sld = root.find(qn(P, "cSld"))
    bg = c_sld.find(qn(P, "bg"))
    if bg is not None:
        c_sld.remove(bg)
    bg = etree.Element(qn(P, "bg"))
    c_sld.insert(0, bg)
    bg_pr = etree.SubElement(bg, qn(P, "bgPr"))
    fill = etree.SubElement(bg_pr, qn(A, "solidFill"))
    etree.SubElement(fill, qn(A, "srgbClr")).set("val", color)
    etree.SubElement(bg_pr, qn(A, "effectLst"))


def _update_presentation(files: dict, count: int) -> None:
    prs = etree.fromstring(files["ppt/presentation.xml"])
    rels = etree.fromstring(files["ppt/_rels/presentation.xml.rels"])

    for rel in list(rels):
        if rel.get("Type") == f"{R_NS}/slide":
            rels.remove(rel)
    used = [int(r.get("Id")[3:]) for r in rels if r.get("Id", "")[3:].isdigit()]
    next_rid = max(used, default=0) + 1

    sld_id_lst = prs.find(qn(P, "sldIdLst"))
    for child in list(sld_id_lst):
        sld_id_lst.remove(child)
    # Произвольные показы и разделы шаблона ссылаются на удалённые слайды
    for el in prs.findall(qn(P, "custShowLst")):
        prs.remove(el)
    for ext in prs.findall(f"{qn(P, 'extLst')}/{qn(P, 'ext')}"):
        if any(etree.QName(child).localname == "sectionLst" for child in ext):
            ext.getparent().remove(ext)

    for i in range(1, count + 1):
        rid = f"rId{next_rid + i - 1}"
        rel = etree.SubElement(rels, qn(PKG_REL, "Relationship"))
        rel.set("Id", rid)
        rel.set("Type", f"{R_NS}/slide")
        rel.set("Target", f"slides/slide{i}.xml")
        sld = etree.SubElement(sld_id_lst, qn(P, "sldId"))
        sld.set("id", str(255 + i))
        sld.set(qn(R_NS, "id"), rid)

    files["ppt/presentation.xml"] = _xml(prs)
    files["ppt/_rels/presentation.xml.rels"] = _xml(rels)


def _update_content_types(files: dict, count: int, image_exts: set[str]) -> None:
    root = etree.fromstring(files["[Content_Types].xml"])
    for el in list(root):
        part = el.get("PartName", "").lstrip("/")
        if el.get("ContentType") == SLIDE_CT or (part and part not in files):
            root.remove(el)
    defaults = {el.get("Extension", "").lower() for el in root.findall(qn(CT_NS, "Default"))}
    for ext in sorted(image_exts | {"png"}):
        if ext not in defaults:
            d = etree.Element(qn(CT_NS, "Default"))
            d.set("Extension", ext)
            d.set("ContentType", IMAGE_CT[ext])
            root.insert(0, d)
    for i in range(1, count + 1):
        o = etree.SubElement(root, qn(CT_NS, "Override"))
        o.set("PartName", f"/ppt/slides/slide{i}.xml")
        o.set("ContentType", SLIDE_CT)
    files["[Content_Types].xml"] = _xml(root)


def _update_app_props(files: dict, count: int) -> None:
    app = files.get("docProps/app.xml")
    if app:
        files["docProps/app.xml"] = re.sub(rb"<Slides>\d+</Slides>", f"<Slides>{count}</Slides>".encode(), app)


# ─── XML хелперы ─────────────────────────────────────────────────────────────

def _xml(root) -> bytes:
    return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def _empty_rels() -> bytes:
    return f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Relationships xmlns="{PKG_REL}"/>'.encode()


def _rels_path(part: str) -> str:
    folder, name = part.rsplit("/", 1)
    return f"{folder}/_rels/{name}.rels"


def _rels_map(files: dict, part: str) -> dict[str, tuple[str, str]]:
    """rId → (тип, цель) для части пакета."""
    data = files.get(_rels_path(part))
    if not data:
        return {}
    return {rel.get("Id"): (rel.get("Type", ""), rel.get("Target", "")) for rel in etree.fromstring(data)}


def _resolve(part: str, target: str) -> str:
    """Относительный путь из .rels → путь внутри архива."""
    if target.startswith("/"):
        return target.lstrip("/")
    parts = part.split("/")[:-1]
    for piece in target.split("/"):
        if piece == "..":
            parts.pop()
        elif piece and piece != ".":
            parts.append(piece)
    return "/".join(parts)


def _shape_id(sp) -> str:
    return sp.find(f"{qn(P, 'nvSpPr')}/{qn(P, 'cNvPr')}").get("id")


def _make_run(text: str, bold: bool = False, sz: int = 3200, color: str | None = None):
    r = etree.Element(qn(A, "r"))
    r_pr = etree.SubElement(r, qn(A, "rPr"))
    r_pr.set("lang", "ru-RU")
    r_pr.set("sz", str(sz))
    r_pr.set("b", "1" if bold else "0")
    r_pr.set("i", "0")
    r_pr.set("u", "none")
    r_pr.set("strike", "noStrike")
    r_pr.set("cap", "none")
    if color:
        fill = etree.SubElement(r_pr, qn(A, "solidFill"))
        etree.SubElement(fill, qn(A, "srgbClr")).set("val", color)
    etree.SubElement(r_pr, qn(A, "latin")).set("typeface", "Calibri")
    etree.SubElement(r_pr, qn(A, "cs")).set("typeface", "Calibri")
    etree.SubElement(r, qn(A, "t")).text = text
    return r


def _new_paragraph(tf, algn: str = "just"):
    p = etree.SubElement(tf, qn(A, "p"))
    p_pr = etree.SubElement(p, qn(A, "pPr"))
    p_pr.set("marL", "0")
    p_pr.set("indent", "0")
    p_pr.set("algn", algn)
    ln_spc = etree.SubElement(p_pr, qn(A, "lnSpc"))
    etree.SubElement(ln_spc, qn(A, "spcPct")).set("val", "90000")
    # Отступы между абзацами обнуляем явно, иначе они наследуются из мастер-слайда
    for tag in ("spcBef", "spcAft"):
        etree.SubElement(etree.SubElement(p_pr, qn(A, tag)), qn(A, "spcPts")).set("val", "0")
    etree.SubElement(p_pr, qn(A, "buNone"))
    return p


def _clear(tf) -> None:
    for p in tf.findall(qn(A, "p")):
        tf.remove(p)
