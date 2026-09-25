"""
Умный парсер пакетов вопросов/ответов.

Логика:
1. Документ превращается в поток строк: абзацы и ячейки таблиц в порядке
   следования, мягкие переносы (Shift+Enter) режутся на отдельные строки,
   картинки запоминаются рядом со строкой, в которой стоят.
2. Конечный автомат раскладывает строки по полям: вопрос, ответ, зачёт,
   комментарий. Поддерживаются «Вопрос 5:», «5.», «5)», автонумерация Word,
   ответ на следующей строке, многострочные вопросы и комментарии, туры.
3. Если есть признаки пропусков (дыры в нумерации, лишние «Ответ:»),
   Scout (Llama 4) ищет пропущенные вопросы. Модель возвращает только
   «якоря» (первые и последние слова), текст берётся из файла дословно.
"""

import difflib
import json
import logging
import re
import tempfile
from dataclasses import dataclass, field
from pathlib import Path

import httpx
from docx import Document
from docx.image.image import Image as DocxImage
from docx.table import Table
from docx.text.paragraph import Paragraph

import office
from config import GROQ_API_KEY, GROQ_BASE_URL, GROQ_MODEL

logger = logging.getLogger(__name__)

# ─── Паттерны ────────────────────────────────────────────────────────────────
# Метка должна заканчиваться разделителем или концом строки,
# иначе «Ответственность…» или «Автор романа…» примем за служебные строки
_SEP = r"(?:\s*[:.\-–—)]\s*|\s*$)"

RE_QUESTION = re.compile(r"^(?:Вопрос|Задание)\s*№?\s*(\d{1,3})\s*[:.)\-–—]?\s*(.*)$", re.I | re.S)
RE_QUESTION_BARE = re.compile(r"^Вопрос\s*[:.]\s*(.*)$", re.I | re.S)
RE_QUESTION_NUM = re.compile(r"^(\d{1,3})\s*[.)](?!\d)\s*(.+)$", re.S)
RE_ANSWER = re.compile(rf"^Ответы?{_SEP}(.*)$", re.I | re.S)
RE_ACCEPT = re.compile(rf"^Зач[её]т{_SEP}(.*)$", re.I | re.S)
RE_REJECT = re.compile(rf"^Незач[её]т{_SEP}(.*)$", re.I | re.S)
RE_COMMENT = re.compile(rf"^Ком+ентари[йи]{_SEP}(.*)$", re.I | re.S)
RE_SERVICE = re.compile(
    r"^(Источник|Источники|Автор|Авторы|Редактор|Редакторы|Ссылка|Ссылки|Апелляция)"
    r"(?:\s*\([^)]*\))?\s*[:.]",
    re.I,
)
RE_SKIP = re.compile(r"^(Давать|Зайти|Примечание)", re.I)  # инструкции ведущему
RE_TOUR = re.compile(r"^(?:Тур\s*№?\s*(\d+)|(\d+)\s*[-–]?\s*(?:й|ый|ой)?\s*тур)(?!\w)\.?", re.I)
RE_HEADER = re.compile(
    r"^(Синхронный|Окский|Открытый|Чемпионат|Турнир|\d{4}-\d{2}-\d{2}|www\.|http|"
    r"С\s+Н2О|С\s+новым)",
    re.I,
)
RE_INLINE_MARKER = re.compile(r"\s+(?=(?:Ответ|Зач[её]т|Комментарий)\s*:)", re.I)

HARD_MARKS = ("*", "★")
ANCHOR_WORDS = 5          # сколько слов берём с начала и конца для якоря
SCOUT_MAX_CHARS = 12000   # сколько текста документа отдаём Scout
SENTENCE_END = tuple(".!?:;…»\")")

# Картинки, которые умеет вставить PowerPoint без конвертации
IMAGE_TYPES = {"image/png", "image/jpeg", "image/gif", "image/bmp"}
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


# ─── Модели ──────────────────────────────────────────────────────────────────

@dataclass
class Picture:
    blob: bytes
    ext: str
    content_type: str
    width: int   # px
    height: int  # px


@dataclass
class Line:
    text: str
    numbered: bool = False              # абзац из автонумерованного списка Word
    pictures: list[Picture] = field(default_factory=list)


@dataclass
class ParseResult:
    questions: list[dict]
    skipped: list[str]                  # что нашли, но не смогли взять (нет ответа и т.п.)
    tours: list[str]                    # названия туров в порядке появления
    ai_checked: bool = False            # Scout реально отработал
    ai_added: int = 0                   # сколько вопросов добавил Scout

    @property
    def pictures(self) -> int:
        return sum(len(q["q_pictures"]) + len(q["a_pictures"]) for q in self.questions)


# ─── Основная функция ────────────────────────────────────────────────────────

def parse_file(path: str) -> ParseResult:
    """Парсит .docx/.doc/.rtf/.odt/.txt. Блокирующая — вызывать через to_thread."""
    ext = Path(path).suffix.lower()
    if ext == ".txt":
        raw = Path(path).read_bytes()
        return parse_text(_decode(raw))
    if ext in (".doc", ".rtf", ".odt"):
        with tempfile.TemporaryDirectory() as tmp:
            converted = office.convert(path, "docx", tmp)
            return _parse_docx(converted)
    return _parse_docx(path)


def parse_text(text: str) -> ParseResult:
    """Парсит обычный текст (из .txt или вставленный в чат)."""
    lines = _split_inline_markers([Line(clean(s)) for s in text.splitlines()])
    return _finalize(_parse_lines(lines), lines)


def looks_like_questions(text: str) -> bool:
    """Похоже ли сообщение на пакет вопросов (а не на обычную фразу)."""
    has_answer = any(RE_ANSWER.match(clean(s)) for s in text.splitlines())
    has_question = any(
        RE_QUESTION.match(clean(s)) or RE_QUESTION_NUM.match(clean(s)) for s in text.splitlines()
    )
    return has_answer and has_question


def _parse_docx(path: str) -> ParseResult:
    lines = _split_inline_markers(read_docx_lines(path))
    return _finalize(_parse_lines(lines), lines)


def _finalize(parsed: tuple[list[dict], list[str], list[str]], lines: list[Line]) -> ParseResult:
    questions, skipped, tours = parsed
    result = ParseResult(questions, skipped, tours)

    if _needs_scout(lines, questions):
        extra = _scout_find_missing(lines, questions)
        result.ai_checked = extra is not None
        if extra:
            questions.extend(extra)
            questions.sort(key=lambda q: q["pos"])
            result.ai_added = len(extra)

    for i, q in enumerate(questions, 1):
        q["number"] = i
    return result


# ─── Чтение документа ────────────────────────────────────────────────────────

def read_docx_lines(path: str) -> list[Line]:
    """Абзацы и ячейки таблиц в порядке следования, с картинками."""
    doc = Document(path)
    formats = _numbering_formats(doc)
    lines: list[Line] = []

    def add_paragraph(p: Paragraph) -> None:
        pictures = _paragraph_pictures(p)
        numbered = _is_numbered(p, formats)
        parts = [clean(s) for s in p.text.split("\n")]
        parts = [s for s in parts if s] or ([""] if pictures else [])
        for i, part in enumerate(parts):
            lines.append(Line(part, numbered and i == 0, pictures if i == len(parts) - 1 else []))

    for el in doc.element.body.iterchildren():
        tag = el.tag.rsplit("}", 1)[-1]
        if tag == "p":
            add_paragraph(Paragraph(el, doc))
        elif tag == "tbl":
            seen = set()
            for row in Table(el, doc).rows:
                for cell in row.cells:
                    if id(cell._tc) in seen:  # объединённые ячейки повторяются
                        continue
                    seen.add(id(cell._tc))
                    for p in cell.paragraphs:
                        add_paragraph(p)
    return lines


def _numbering_formats(doc) -> dict[str, dict[str, str]]:
    """numId → {уровень: формат} ("decimal", "bullet", …) из numbering.xml."""
    try:
        numbering = doc.part.numbering_part.element
    except Exception:
        return {}
    abstract = {}
    for an in numbering.findall(f"{{{W_NS}}}abstractNum"):
        levels = {}
        for lvl in an.findall(f"{{{W_NS}}}lvl"):
            fmt = lvl.find(f"{{{W_NS}}}numFmt")
            levels[lvl.get(f"{{{W_NS}}}ilvl")] = fmt.get(f"{{{W_NS}}}val") if fmt is not None else "decimal"
        abstract[an.get(f"{{{W_NS}}}abstractNumId")] = levels
    formats = {}
    for num in numbering.findall(f"{{{W_NS}}}num"):
        ref = num.find(f"{{{W_NS}}}abstractNumId")
        if ref is not None:
            formats[num.get(f"{{{W_NS}}}numId")] = abstract.get(ref.get(f"{{{W_NS}}}val"), {})
    return formats


def _is_numbered(p: Paragraph, formats: dict) -> bool:
    """Абзац из нумерованного (не маркированного) списка Word. Номер не входит в p.text."""
    try:
        sources = [p._p.pPr]
        style = p.style
        while style is not None:
            sources.append(style.element.pPr)
            style = style.base_style
        for pPr in sources:
            num_pr = pPr.find(f"{{{W_NS}}}numPr") if pPr is not None else None
            if num_pr is None:
                continue
            num_id = num_pr.find(f"{{{W_NS}}}numId")
            ilvl = num_pr.find(f"{{{W_NS}}}ilvl")
            num_id = num_id.get(f"{{{W_NS}}}val") if num_id is not None else None
            level = ilvl.get(f"{{{W_NS}}}val") if ilvl is not None else "0"
            if not num_id or num_id == "0":
                return False
            return formats.get(num_id, {}).get(level, "decimal") not in ("bullet", "none")
    except Exception:
        pass
    return False


def _paragraph_pictures(p: Paragraph) -> list[Picture]:
    """Вытаскивает встроенные картинки из абзаца (DrawingML и старый VML)."""
    pictures = []
    for el in p._p.iter():
        if not isinstance(el.tag, str):
            continue
        tag = el.tag.rsplit("}", 1)[-1]
        if tag == "blip":
            rid = el.get(f"{{{R_NS}}}embed")
        elif tag == "imagedata":
            rid = el.get(f"{{{R_NS}}}id")
        else:
            continue
        if not rid:
            continue
        try:
            part = p.part.related_parts[rid]
            img = DocxImage.from_blob(part.blob)
            if img.content_type not in IMAGE_TYPES:
                continue
            pictures.append(Picture(part.blob, img.ext, img.content_type, img.px_width, img.px_height))
        except Exception as e:
            logger.info("Картинка пропущена: %s", e)
    return pictures


# ─── Конечный автомат ────────────────────────────────────────────────────────

def _split_inline_markers(lines: list[Line]) -> list[Line]:
    """«Столица Франции? Ответ: Париж» → две строки."""
    result = []
    for line in lines:
        parts = RE_INLINE_MARKER.split(line.text)
        if len(parts) == 1:
            result.append(line)
            continue
        for i, part in enumerate(parts):
            last = i == len(parts) - 1
            result.append(Line(part, line.numbered and i == 0, line.pictures if last else []))
    return result


def _parse_lines(lines: list[Line]) -> tuple[list[dict], list[str], list[str]]:
    questions: list[dict] = []
    skipped: list[str] = []
    tours: list[str] = []
    current: dict | None = None
    field_name: str | None = None   # куда дописывать продолжение
    tour: str | None = None

    def finish() -> None:
        if current is None:
            return
        label = f"№{current['orig_number']}" if current["orig_number"] else _preview(current["question"])
        if not current["answer"]:
            skipped.append(f"{label} — нет ответа")
        elif not current["question"] and not current["q_pictures"]:
            skipped.append(f"{label} — пустой текст вопроса")
        else:
            questions.append(current)

    for pos, line in enumerate(lines):
        text = line.text
        if not text and not line.pictures:
            continue

        # Тур («Тур 2», «2 тур») — короткая отдельная строка
        m_tour = RE_TOUR.match(text) if len(text) <= 30 else None
        if m_tour:
            finish()
            current, field_name = None, None
            tour = f"Тур {m_tour.group(1) or m_tour.group(2)}"
            if tour not in tours:
                tours.append(tour)
            continue

        if RE_SKIP.match(text):
            continue

        between = current is None or bool(current["answer"])
        start = _question_start(line, between)
        if start is not None:
            orig, q_text = start
            if orig is None and line.numbered:
                # Автонумерация Word: номер в тексте не виден, считаем сами
                orig = (current["orig_number"] or 0) + 1 if current else 1
            finish()
            hard = q_text.startswith(HARD_MARKS)
            if hard:
                q_text = q_text.lstrip("*★ ").strip()
            current = {
                "number": 0,
                "orig_number": orig,
                "question": q_text,
                "answer": "",
                "accept": None,
                "comment": None,
                "hard": hard,
                "tour": tour,
                "pos": pos,
                "q_pictures": list(line.pictures),
                "a_pictures": [],
            }
            field_name = "question"
            continue

        if current is None:
            continue

        if m := RE_ANSWER.match(text):
            if not current["answer"]:
                current["answer"] = m.group(1).strip()
                field_name = "answer"
            else:
                field_name = None
            current["a_pictures"] += line.pictures
            continue
        if m := RE_ACCEPT.match(text):
            current["accept"] = m.group(1).strip()
            field_name = "accept"
            continue
        if RE_REJECT.match(text):
            field_name = None
            continue
        if m := RE_COMMENT.match(text):
            current["comment"] = m.group(1).strip()
            field_name = "comment"
            current["a_pictures"] += line.pictures
            continue
        if RE_SERVICE.match(text) or (field_name != "question" and RE_HEADER.match(text)):
            field_name = None
            continue

        # Продолжение текущего поля (строка может быть пустой, если в ней только картинка)
        if field_name == "question":
            if not current["question"] and text.startswith(HARD_MARKS):
                current["hard"] = True
                text = text.lstrip("*★ ").strip()
            if text:
                current["question"] = _join(current["question"], text)
            current["q_pictures"] += line.pictures
        elif field_name == "answer" and not current["answer"]:
            current["answer"] = text  # «Ответ:» на отдельной строке
            current["a_pictures"] += line.pictures
        elif field_name in ("accept", "comment"):
            if text:
                current[field_name] = _join(current[field_name] or "", text)
            current["a_pictures"] += line.pictures
        elif field_name == "answer":
            current["a_pictures"] += line.pictures

    finish()
    return questions, skipped, tours


def _question_start(line: Line, between: bool) -> tuple[int | None, str] | None:
    """Если строка начинает новый вопрос — (номер, текст), иначе None."""
    text = line.text
    if m := RE_QUESTION.match(text):
        return int(m.group(1)), m.group(2).strip()
    if m := RE_QUESTION_BARE.match(text):
        return None, m.group(1).strip()
    # «5. Текст» и автонумерация Word — только между вопросами, иначе это
    # список внутри вопроса или год в начале строки
    if not between or _is_marker(text):
        return None
    if m := RE_QUESTION_NUM.match(text):
        return int(m.group(1)), m.group(2).strip()
    if line.numbered and len(text) > 15:
        return None, text
    return None


def _is_marker(text: str) -> bool:
    return any(r.match(text) for r in (RE_ANSWER, RE_ACCEPT, RE_REJECT, RE_COMMENT, RE_SERVICE))


def _join(prev: str, text: str) -> str:
    """Склеивает строки: абзацы — через перенос, обрывки фраз — через пробел."""
    if not prev:
        return text
    if prev.endswith(SENTENCE_END) or not text[:1].islower():
        return f"{prev}\n{text}"
    return f"{prev} {text}"


# ─── Scout с якорным поиском ─────────────────────────────────────────────────

def _needs_scout(lines: list[Line], questions: list[dict]) -> bool:
    """Звать ИИ только если есть признаки пропусков — экономим время и лимиты."""
    if not GROQ_API_KEY:
        return False
    if not questions:
        return True
    answers = sum(1 for ln in lines if RE_ANSWER.match(ln.text))
    if answers > len(questions):
        return True
    nums = [q["orig_number"] for q in questions if q["orig_number"]]
    for a, b in zip(nums, nums[1:]):
        if b - a > 1:  # дыра в нумерации (сброс нумерации в новом туре — не дыра)
            return True
    return False


def _scout_find_missing(lines: list[Line], existing: list[dict]) -> list[dict] | None:
    """Scout возвращает якоря пропущенных вопросов, текст берём из файла.
    None — если ИИ недоступен (проверка не состоялась)."""
    paragraphs = [ln.text for ln in lines if ln.text]
    full_text = "\n".join(paragraphs)[:SCOUT_MAX_CHARS]
    existing_summary = json.dumps(
        [{"n": q["orig_number"] or i + 1, "q_start": _words(q["question"], ANCHOR_WORDS, "start")}
         for i, q in enumerate(existing)],
        ensure_ascii=False,
    )

    prompt = f"""Ты проверяешь полноту списка вопросов Brain-Do.

УЖЕ НАЙДЕННЫЕ вопросы (начало каждого):
{existing_summary}

ПОЛНЫЙ ТЕКСТ ДОКУМЕНТА:
{full_text}

ЗАДАЧА: найди вопросы которых НЕТ в списке выше.

Для каждого ПРОПУЩЕННОГО вопроса верни якоря — первые {ANCHOR_WORDS} и последние {ANCHOR_WORDS} слов вопроса и ответа.
Если пропущенных нет — верни пустой массив [].

СТРОГО верни ТОЛЬКО JSON без пояснений:
[
  {{
    "q_start": "первые {ANCHOR_WORDS} слов вопроса",
    "q_end": "последние {ANCHOR_WORDS} слов вопроса",
    "a_start": "первые {ANCHOR_WORDS} слов ответа",
    "a_end": "последние {ANCHOR_WORDS} слов ответа"
  }}
]"""

    try:
        response = httpx.post(
            GROQ_BASE_URL,
            headers={"Authorization": f"Bearer {GROQ_API_KEY}"},
            json={
                "model": GROQ_MODEL,
                "messages": [{"role": "user", "content": prompt}],
                "temperature": 0.0,
                "max_tokens": 2000,
            },
            timeout=40,
        )
        response.raise_for_status()
        raw = response.json()["choices"][0]["message"]["content"]
        raw = re.sub(r"```json|```", "", raw).strip()
        anchors = json.loads(raw) or []
    except Exception as e:
        logger.warning("Scout недоступен: %s", e)
        return None

    known = {_normalize(q["question"])[:60] for q in existing}
    extra = []
    for anchor in anchors:
        if not isinstance(anchor, dict):
            continue
        q_pos, q_text = _find_by_anchors(paragraphs, anchor.get("q_start", ""), anchor.get("q_end", ""))
        _, a_text = _find_by_anchors(paragraphs, anchor.get("a_start", ""), anchor.get("a_end", ""))
        if not (q_text and a_text) or _normalize(q_text)[:60] in known:
            continue
        known.add(_normalize(q_text)[:60])
        extra.append({
            "number": 0, "orig_number": None, "question": q_text, "answer": a_text,
            "accept": None, "comment": None, "hard": False, "tour": None,
            "pos": _line_pos(lines, q_pos, paragraphs), "q_pictures": [], "a_pictures": [],
        })
    return extra


def _line_pos(lines: list[Line], paragraph_idx: int, paragraphs: list[str]) -> float:
    """Позиция найденного абзаца в исходном потоке строк (для сортировки)."""
    target = paragraphs[paragraph_idx] if 0 <= paragraph_idx < len(paragraphs) else None
    for i, ln in enumerate(lines):
        if ln.text == target:
            return i + 0.5
    return float(len(lines))


def _find_by_anchors(paragraphs: list[str], start_anchor: str, end_anchor: str) -> tuple[int, str]:
    """Ищет абзац(ы) по якорям начала и конца. Возвращает (индекс, дословный текст)."""
    if not start_anchor:
        return -1, ""

    start_words = _normalize(start_anchor)
    end_words = _normalize(end_anchor or "")
    best_score, best_idx, best_text = 0.0, -1, ""

    for i, p in enumerate(paragraphs):
        for candidate in (p, p + " " + paragraphs[i + 1] if i + 1 < len(paragraphs) else None):
            if not candidate:
                continue
            norm = _normalize(candidate)
            if not norm:
                continue
            score_start = difflib.SequenceMatcher(None, start_words[:40], norm[:40]).ratio()
            score_end = (
                difflib.SequenceMatcher(None, end_words[-40:], norm[-40:]).ratio() if end_words else 0.5
            )
            score = score_start * 0.6 + score_end * 0.4
            if score > best_score and score > 0.4:
                best_score, best_idx, best_text = score, i, candidate

    best_text = re.sub(r"^(Ответ|Комментарий|Зач[её]т)\s*[:.]?\s*", "", best_text, flags=re.I)
    best_text = re.sub(r"^(Вопрос\s*\d*|\d{1,3})\s*[:.)]\s*", "", best_text, flags=re.I)
    return best_idx, clean(best_text)


# ─── Утилиты ─────────────────────────────────────────────────────────────────

def _decode(raw: bytes) -> str:
    for enc in ("utf-8-sig", "cp1251"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


def _normalize(text: str) -> str:
    """Нормализуем для сравнения — нижний регистр, без пунктуации."""
    text = text.lower().replace("\xa0", " ")
    text = re.sub(r"[^\w\s]", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def _words(text: str, n: int, side: str) -> str:
    words = text.split()
    return " ".join(words[:n] if side == "start" else words[-n:])


def _preview(text: str, n: int = 40) -> str:
    text = text.replace("\n", " ")
    return f"«{text[:n]}…»" if len(text) > n else f"«{text}»"


def clean(text: str) -> str:
    text = text.replace("\xa0", " ").replace("\t", " ")
    text = re.sub(r" {2,}", " ", text)
    return text.strip()
